"""
Resume ingestion and storage.

Purpose
-------
Turns an uploaded PDF/DOCX/TXT CV into the plain text that bounds every AI
claim, and manages which variant is the default.

How it works
------------
The file is written to ``RESUME_DIR`` under a sanitised name, its text is
extracted, and both are stored. ``content_text`` is what the AI layer and the
fabrication guard read — the binary is kept only so it can be re-uploaded to
LinkedIn.

Failure scenarios
-----------------
* **Scanned/image-only PDF** → extraction yields almost no text.
  :class:`ResumeExtractionError` is raised with an explicit instruction to run
  OCR or paste the text, because silently storing an empty resume would make
  every score 0 and every letter unsupportable.
* **Encrypted PDF** → :class:`ResumeExtractionError` naming the problem.
* **Unsupported extension** → :class:`ResumeExtractionError` listing what is
  supported.
"""

from __future__ import annotations

import io
import pathlib
import re
import unicodedata

from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import Resume
from app.logging_config import get_logger

log = get_logger(__name__)

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".md"})

# Below this many characters, a "successful" extraction is almost certainly a
# scanned document rather than a real CV.
_MIN_PLAUSIBLE_LENGTH = 200


class ResumeExtractionError(RuntimeError):
    """Raised when a resume file cannot be turned into usable text."""


def _sanitise_filename(name: str) -> str:
    """Reduce an uploaded filename to something safe to write to disk."""
    stem = pathlib.Path(name).name
    stem = unicodedata.normalize("NFKD", stem).encode("ascii", "ignore").decode()
    stem = re.sub(r"[^A-Za-z0-9._-]", "_", stem).strip("._")
    return stem or "resume"


def extract_text(data: bytes, filename: str) -> str:
    """Extract plain text from a resume file.

    Args:
        data: Raw file bytes.
        filename: Original filename; its extension selects the parser.

    Raises:
        ResumeExtractionError: unsupported type, unreadable file, or too little
            text to be a real CV.
    """
    suffix = pathlib.Path(filename).suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeExtractionError(
            f"Unsupported resume type {suffix!r}. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix in {".txt", ".md"}:
        text = data.decode("utf-8", errors="replace")
    elif suffix == ".pdf":
        text = _extract_pdf(data)
    else:
        text = _extract_docx(data)

    text = _normalise_whitespace(text)

    if len(text) < _MIN_PLAUSIBLE_LENGTH:
        raise ResumeExtractionError(
            f"Only {len(text)} characters of text were extracted from {filename!r}. "
            "This usually means the file is a scanned image rather than a text PDF. "
            "Run it through OCR, export a text-based PDF, or upload a .txt version. "
            "Storing an empty resume would make every score meaningless."
        )

    return text


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise ResumeExtractionError(f"Could not read PDF: {exc}") from exc

    if reader.is_encrypted:
        # An empty user password is common for "protected" CVs and decrypts fine.
        try:
            if reader.decrypt("") == 0:
                raise ResumeExtractionError(
                    "PDF is password-protected. Remove the password and re-upload."
                )
        except Exception as exc:
            raise ResumeExtractionError(
                "PDF is encrypted and could not be opened. Remove the password and re-upload."
            ) from exc

    parts: list[str] = []
    for index, page in enumerate(reader.pages):
        try:
            parts.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 — one bad page must not lose the rest
            log.warning("resume.pdf_page_failed", page=index)
    return "\n".join(parts)


def _extract_docx(data: bytes) -> str:
    import docx  # python-docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeExtractionError(f"Could not read DOCX: {exc}") from exc

    parts = [p.text for p in document.paragraphs]
    # Skills and dates are very often in tables, which paragraphs miss entirely.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _normalise_whitespace(text: str) -> str:
    """Collapse the ragged whitespace PDF extraction produces."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def save_file(data: bytes, filename: str, resume_dir: str) -> str:
    """Persist the original file and return its path."""
    directory = pathlib.Path(resume_dir)
    directory.mkdir(parents=True, exist_ok=True)

    safe = _sanitise_filename(filename)
    target = directory / safe

    # Never silently overwrite an existing CV.
    counter = 1
    while target.exists():
        stem = pathlib.Path(safe).stem
        suffix = pathlib.Path(safe).suffix
        target = directory / f"{stem}_{counter}{suffix}"
        counter += 1

    target.write_bytes(data)
    target.chmod(0o600)
    log.info("resume.file_saved", path=str(target), bytes=len(data))
    return str(target)


async def set_default_resume(session: AsyncSession, resume_id: int) -> None:
    """Make ``resume_id`` the default, clearing any previous one.

    Clearing first is required: a partial unique index enforces at most one
    default, so setting the new one before clearing the old would violate it.
    """
    await session.execute(
        update(Resume).where(Resume.is_default.is_(True)).values(is_default=False)
    )
    await session.execute(update(Resume).where(Resume.id == resume_id).values(is_default=True))
    log.info("resume.default_changed", resume_id=resume_id)


async def get_default_resume(session: AsyncSession) -> Resume | None:
    """Return the default resume, falling back to any active one."""
    stmt = (
        select(Resume)
        .where(Resume.is_active.is_(True))
        .order_by(Resume.is_default.desc(), Resume.id.asc())
        .limit(1)
    )
    return (await session.execute(stmt)).scalar_one_or_none()


def guess_skills(text: str, *, limit: int = 40) -> list[str]:
    """Best-effort skill extraction from resume text.

    Purely a convenience for the dashboard — nothing decision-making depends on
    it. Scoring uses the full resume text, not this list, precisely so that a
    missed skill here can never cost a match.
    """
    known = {
        "python",
        "javascript",
        "typescript",
        "java",
        "go",
        "rust",
        "ruby",
        "php",
        "c++",
        "c#",
        "kotlin",
        "swift",
        "scala",
        "sql",
        "bash",
        "r",
        "matlab",
        "react",
        "vue",
        "angular",
        "svelte",
        "next.js",
        "node.js",
        "express",
        "django",
        "flask",
        "fastapi",
        "spring",
        "rails",
        "laravel",
        ".net",
        "postgresql",
        "postgres",
        "mysql",
        "mongodb",
        "redis",
        "elasticsearch",
        "sqlite",
        "cassandra",
        "dynamodb",
        "neo4j",
        "aws",
        "azure",
        "gcp",
        "docker",
        "kubernetes",
        "terraform",
        "ansible",
        "jenkins",
        "github actions",
        "gitlab ci",
        "circleci",
        "argocd",
        "helm",
        "pytorch",
        "tensorflow",
        "scikit-learn",
        "pandas",
        "numpy",
        "spark",
        "airflow",
        "dbt",
        "kafka",
        "rabbitmq",
        "graphql",
        "rest",
        "grpc",
        "playwright",
        "selenium",
        "cypress",
        "pytest",
        "jest",
        "junit",
        "linux",
        "git",
        "nginx",
        "prometheus",
        "grafana",
        "datadog",
        "elk",
    }
    lowered = text.lower()
    found = sorted({skill for skill in known if skill in lowered})
    return found[:limit]
